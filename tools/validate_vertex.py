#!/usr/bin/env python3
"""Validate the reconstructed Vertex engine against a TradingView Excel export.

The computation is intentionally dependency-free. Reading .xlsx files uses
openpyxl, imported only when needed. CSV files use the Python standard library.
"""

from __future__ import annotations

import argparse
import csv
import math
from dataclasses import dataclass
from pathlib import Path
from statistics import fmean, pstdev
from typing import Any, Iterable, Mapping, Sequence


EXPECTED_COLUMNS = {
    "control": "control period",
    "signal": "signal period",
    "upper": "upper bb of vertex",
    "lower": "lower bb of vertex",
}
MISSING_TEXT = {"", "∅", "na", "n/a", "nan", "none", "null"}


@dataclass(frozen=True)
class VertexSettings:
    control_period: int = 14
    signal_period: int = 5
    upper_period: int = 12
    upper_deviation: float = 2.0
    lower_period: int = 12
    lower_deviation: float = 2.0
    use_previous_bar: bool = True


@dataclass(frozen=True)
class VertexSeries:
    control: list[float]
    signal: list[float | None]
    upper: list[float | None]
    lower: list[float | None]


def _validate_settings(settings: VertexSettings) -> None:
    lengths = (
        settings.control_period,
        settings.signal_period,
        settings.upper_period,
        settings.lower_period,
    )
    if any(length < 1 for length in lengths):
        raise ValueError("All periods must be at least 1")
    if settings.upper_deviation < 0 or settings.lower_deviation < 0:
        raise ValueError("Band deviations cannot be negative")


def vertex_control(
    highs: Sequence[float],
    lows: Sequence[float],
    closes: Sequence[float],
    settings: VertexSettings = VertexSettings(),
) -> list[float]:
    """Return the Vertex control series in oldest-to-newest order."""
    _validate_settings(settings)
    if not (len(highs) == len(lows) == len(closes)):
        raise ValueError("High, low, and close series must have equal lengths")

    output: list[float] = []
    delay = 1 if settings.use_previous_bar else 0

    for current in range(len(closes)):
        newest = current - delay
        if newest < 0:
            output.append(0.0)
            continue

        oldest = max(0, newest - settings.control_period + 1)
        running_high: float | None = None
        running_low: float | None = None
        up_sum = 0.0
        down_sum = 0.0

        for index in range(newest, oldest - 1, -1):
            high = float(highs[index])
            low = float(lows[index])
            close = float(closes[index])
            if not all(math.isfinite(value) for value in (high, low, close)):
                raise ValueError(f"Non-finite OHLC value at chronological row {index + 1}")

            if running_high is None or high > running_high:
                running_high = high
                up_sum += close
            if running_low is None or low < running_low:
                running_low = low
                down_sum += close

        if up_sum == 0.0 or down_sum == 0.0:
            output.append(math.nan)
        else:
            output.append(down_sum / up_sum - up_sum / down_sum)

    return output


def rolling_sma(values: Sequence[float], length: int) -> list[float | None]:
    output: list[float | None] = [None] * len(values)
    for index in range(length - 1, len(values)):
        window = values[index - length + 1 : index + 1]
        if all(math.isfinite(value) for value in window):
            output[index] = fmean(window)
    return output


def rolling_population_band(
    values: Sequence[float], length: int, deviation: float, upper: bool
) -> list[float | None]:
    output: list[float | None] = [None] * len(values)
    direction = 1.0 if upper else -1.0
    for index in range(length - 1, len(values)):
        window = values[index - length + 1 : index + 1]
        if all(math.isfinite(value) for value in window):
            output[index] = fmean(window) + direction * deviation * pstdev(window)
    return output


def calculate_vertex(
    highs: Sequence[float],
    lows: Sequence[float],
    closes: Sequence[float],
    settings: VertexSettings = VertexSettings(),
) -> VertexSeries:
    control = vertex_control(highs, lows, closes, settings)
    return VertexSeries(
        control=control,
        signal=rolling_sma(control, settings.signal_period),
        upper=rolling_population_band(
            control, settings.upper_period, settings.upper_deviation, upper=True
        ),
        lower=rolling_population_band(
            control, settings.lower_period, settings.lower_deviation, upper=False
        ),
    )


def _normalized_row(row: Mapping[str, Any]) -> dict[str, Any]:
    return {str(key).strip().casefold(): value for key, value in row.items()}


def read_rows(path: Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Input must be .xlsx, .xlsm, or .csv")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Reading Excel files requires openpyxl. Run: python -m pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    rows = [dict(zip(headers, values)) for values in iterator]
    workbook.close()
    return rows


def _number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, str):
        cleaned = value.strip().casefold()
        if cleaned in MISSING_TEXT:
            return None
        cleaned = cleaned.replace(",", "")
        try:
            value = float(cleaned)
        except ValueError:
            return None
    number = float(value)
    return number if math.isfinite(number) else None


def _required_number(row: Mapping[str, Any], name: str, row_number: int) -> float:
    value = _number(row.get(name))
    if value is None:
        raise ValueError(f"Missing numeric {name!r} at data row {row_number}")
    return value


def _metrics(
    predicted: Sequence[float | None],
    expected: Sequence[float | None],
    tolerance: float,
) -> dict[str, float | int]:
    errors = [
        abs(float(actual) - float(target))
        for actual, target in zip(predicted, expected)
        if actual is not None
        and target is not None
        and math.isfinite(float(actual))
        and math.isfinite(float(target))
    ]
    if not errors:
        return {"count": 0, "mae": math.nan, "max_error": math.nan, "within": 0, "rate": math.nan}
    within = sum(error <= tolerance for error in errors)
    return {
        "count": len(errors),
        "mae": fmean(errors),
        "max_error": max(errors),
        "within": within,
        "rate": within / len(errors),
    }


def validate_export(
    rows: Iterable[Mapping[str, Any]],
    chronology: str = "descending",
    tolerance: float = 0.00501,
    settings: VertexSettings = VertexSettings(),
) -> dict[str, dict[str, float | int]]:
    normalized = [_normalized_row(row) for row in rows]
    if chronology == "descending":
        chronological = list(reversed(normalized))
    elif chronology == "ascending":
        chronological = normalized
    else:
        raise ValueError("chronology must be 'ascending' or 'descending'")

    highs = [_required_number(row, "high", index + 1) for index, row in enumerate(chronological)]
    lows = [_required_number(row, "low", index + 1) for index, row in enumerate(chronological)]
    closes = [_required_number(row, "close", index + 1) for index, row in enumerate(chronological)]
    calculated = calculate_vertex(highs, lows, closes, settings)

    results: dict[str, dict[str, float | int]] = {}
    for series_name, column_name in EXPECTED_COLUMNS.items():
        expected = [_number(row.get(column_name)) for row in chronological]
        results[series_name] = _metrics(getattr(calculated, series_name), expected, tolerance)
    return results


def _print_report(results: Mapping[str, Mapping[str, float | int]], tolerance: float) -> None:
    print(f"Tolerance: ±{tolerance:.5f}")
    print(f"{'Series':<10} {'Rows':>8} {'MAE':>12} {'Max error':>12} {'Within':>12}")
    for name in ("control", "signal", "upper", "lower"):
        values = results[name]
        print(
            f"{name:<10} {int(values['count']):>8,} "
            f"{float(values['mae']):>12.9f} {float(values['max_error']):>12.9f} "
            f"{float(values['rate']):>11.3%}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("export", type=Path, help="TradingView export (.xlsx, .xlsm, or .csv)")
    parser.add_argument("--sheet", help="Excel sheet name; defaults to the first sheet")
    parser.add_argument(
        "--chronology",
        choices=("descending", "ascending"),
        default="descending",
        help="Row order in the export; the supplied workbook is descending",
    )
    parser.add_argument(
        "--tolerance",
        type=float,
        default=0.00501,
        help="Maximum accepted absolute error; default allows 2-decimal rounding",
    )
    args = parser.parse_args()

    if args.tolerance < 0:
        parser.error("--tolerance cannot be negative")

    results = validate_export(
        read_rows(args.export, args.sheet),
        chronology=args.chronology,
        tolerance=args.tolerance,
    )
    _print_report(results, args.tolerance)
    return 0 if all(float(values["rate"]) == 1.0 for values in results.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())
